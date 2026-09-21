"""네이버 증권(https://stock.naver.com/market/stock/kr) 국내주식 랭킹 크롤러.

두 가지 방식을 지원한다.
1) 기본(API): 화면이 내부적으로 호출하는 JSON 주소를 페이징(startIdx/pageSize)으로 호출한다.
   화면은 자바스크립트로 그려져서 requests + BeautifulSoup으로는 데이터가 비어 있기 때문이다.
2) --html: 브라우저에서 저장한 HTML 파일을 BeautifulSoup4로 파싱한다.
   (거래대금 상위 4개 그룹: 국내주식/해외주식/국내 ETF/가상자산 - GlobalStockRankingV2 영역)

사용법:
    python stock_rank_crawler.py                     # 국내 거래대금 상위 20개
    python stock_rank_crawler.py -n 100              # 100개 (페이징으로 자동 수집)
    python stock_rank_crawler.py -o marketSum -m KOSPI   # 코스피 시가총액순
    python stock_rank_crawler.py -x rank.xlsx -j rank.json
    python stock_rank_crawler.py --html page.html    # 저장된 HTML 파싱 (4개 그룹)
    python stock_rank_crawler.py --html page.html -g crypto
필요 패키지: requests, beautifulsoup4 (엑셀 저장 시 openpyxl)
"""
import argparse
import json
import time

import requests
from bs4 import BeautifulSoup

BASE = "https://stock.naver.com"
API = BASE + "/api/domestic/market/stock/default"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
}
PAGE_SIZE = 50    # 한 번에 요청할 개수
MAX_ITEMS = 1000  # 안전 상한
PAGE_DELAY = 0.3  # 페이지 사이 대기(초)

GROUPS = {  # data-nlogs 값 -> (옵션 이름, 화면 이름)
    "rank.stkkrlist": ("kr", "국내주식"),
    "rank.stkuslist": ("us", "해외주식"),
    "rank.etfkrlist": ("etf", "국내 ETF"),
    "rank.cryptolist": ("crypto", "가상자산"),
}

ORDER_TYPES = {  # API orderType -> 화면 이름
    "priceTop": "거래대금 상위", "marketSum": "시가총액", "quantTop": "거래량 상위",
    "up": "상승", "down": "하락", "searchTop": "인기 종목",
    "high52week": "52주 최고", "low52week": "52주 최저",
    "foreignPureBuy": "외국인 상위", "organizationPureBuy": "기관 상위",
}


# ---------------------------------------------------------------- API 방식
def format_won(amount):
    """원 단위 정수를 '8조 5,320억' 형태로 변환."""
    try:
        eok = int(amount) // 100_000_000
    except (TypeError, ValueError):
        return ""
    jo, eok = divmod(eok, 10_000)
    if jo:
        return f"{jo}조 {eok:,}억" if eok else f"{jo}조"
    return f"{eok:,}억"


def api_item(rank, s):
    code = s["itemcode"]
    up_down = str(s.get("upDownGb", ""))
    direction = {"1": "상승", "2": "상승", "3": "보합", "4": "하락", "5": "하락"}.get(up_down, "")
    sign = {"상승": "+", "하락": "-"}.get(direction, "")
    price = int(s["nowPrice"])
    change = abs(int(s["prevChangePrice"]))  # 하락은 API가 이미 음수로 주므로 절댓값에 부호를 다시 붙임
    rate = abs(float(s["prevChangeRate"]))
    return {
        "rank": rank,
        "name": s["itemname"],
        "code": code,
        "price": f"{price:,}",
        "session": "",
        "direction": direction,
        "change": f"{sign}{change:,}" if change else "0",
        "rate": f"{sign}{rate:.2f}%" if sign else "0.00%",
        "trading_value": format_won(s.get("tradeAmount")),
        "url": f"{BASE}/domestic/stock/{code}/price",
    }


def fetch_domestic(order="priceTop", market="ALL", trade="KRX", limit=20, on_page=None):
    """startIdx를 PAGE_SIZE씩 늘려가며 limit개가 될 때까지(또는 데이터가 끝날 때까지) 수집."""
    limit = min(limit, MAX_ITEMS)
    items, start = [], 0
    while len(items) < limit:
        size = min(PAGE_SIZE, limit - len(items))
        resp = requests.get(API, headers=HEADERS, timeout=10, params={
            "tradeType": trade, "marketType": market, "orderType": order,
            "startIdx": start, "pageSize": size})
        if resp.status_code == 400:
            raise ValueError(f"잘못된 요청 값입니다: {resp.text[:200]}")
        resp.raise_for_status()
        rows = resp.json()
        if not rows:
            break
        items.extend(api_item(start + i + 1, r) for i, r in enumerate(rows))
        if on_page:
            on_page(start // PAGE_SIZE + 1, len(rows), len(items))
        if len(rows) < size:  # 마지막 페이지
            break
        start += len(rows)
        time.sleep(PAGE_DELAY)
    return items


# ---------------------------------------------------------------- HTML 파싱 방식
def read_html_file(path):
    """저장된 HTML 파일을 읽는다. 인코딩(UTF-8/UTF-8 BOM/CP949)을 차례로 시도."""
    try:
        with open(path, "rb") as f:
            raw = f.read()
    except OSError as e:
        raise SystemExit(f"HTML 파일을 읽을 수 없습니다: {e}")
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def text_of(el, selector):
    """CSS 클래스는 빌드마다 뒤 해시(__xxxx)가 바뀌므로 앞부분만 부분 일치로 찾는다."""
    found = el.select_one(selector)
    return " ".join(found.get_text().split()) if found else ""


def parse_item(li):
    a = li.select_one("a[data-nlogs]")
    href = a["href"]

    price = text_of(li, '[class*="ModulePriceNumber_price__"]')
    prefix = text_of(li, '[class*="ModulePriceNumber_prefix__"]')      # 해외주식 '$'
    session = text_of(li, '[class*="HomeTopStockInfo_price-session__"]')  # 해외주식 'Pre' 등

    # 등락: 스크린리더용 '상승/하락' 텍스트(.a11y)를 방향으로, 나머지를 변동폭으로 사용
    change_el = li.select_one('[class*="ModulePriceChange_amount__"]')
    direction = text_of(change_el, ".a11y") if change_el else ""
    change = ""
    if change_el:
        for hidden in change_el.select(".a11y"):
            hidden.extract()
        change = " ".join(change_el.get_text().split())

    trading = li.select_one('[class*="HomeTopStockInfo_trading-value__"]')
    if trading:
        for label in trading.select('[class*="HomeTopStockInfo_label__"]'):
            label.extract()
    rate = text_of(li, '[class*="ModulePercent_module-percent__"]').strip("()")
    sign = {"상승": "+", "하락": "-"}.get(direction, "")

    return {
        "rank": int(text_of(li, '[class*="HomeTopStockInfo_rank-number__"]') or 0),
        "name": text_of(li, '[class*="StockDoubleLineText_text__"]'),
        "code": href.rstrip("/").split("/")[-2] if href.endswith("/price") else href.split("/")[-1],
        "price": prefix + price,
        "session": session,
        "direction": direction,
        "change": (sign + change) if change and not change.startswith(("+", "-")) else change,
        "rate": rate,
        "trading_value": " ".join(trading.get_text().split()) if trading else "",
        "url": BASE + href,
    }


def parse_rankings(html):
    """HTML에서 그룹별 랭킹을 추출. {그룹키: [종목 dict, ...]}"""
    soup = BeautifulSoup(html, "html.parser")
    result = {}
    for link in soup.select('a[data-nlogs^="rank."]'):
        key = link["data-nlogs"]
        if key not in GROUPS:
            continue
        li = link.find_parent("li")
        # 각 그룹에는 무한 스크롤용 복제본 <ol aria-hidden="true">가 있으므로 제외
        if li is None or li.find_parent("ol", attrs={"aria-hidden": "true"}):
            continue
        try:
            result.setdefault(key, []).append(parse_item(li))
        except (AttributeError, KeyError, ValueError, TypeError):
            continue  # 구조가 다른 항목 하나 때문에 전체가 중단되지 않도록 건너뜀
    return result


# ---------------------------------------------------------------- 출력 / 저장
def print_rankings(data, titles):
    for key, items in data.items():
        if not items:
            continue
        print(f"\n[{titles[key]}]  {len(items)}개")
        for it in items:
            sess = f"[{it['session']}] " if it["session"] else ""
            print(f"  {it['rank']:>3}. {it['name']:<16} {sess}{it['price']:>12}  "
                  f"{it['direction']} {it['change']:>9} ({it['rate']})  거래대금 {it['trading_value']}")


def save_excel(data, titles, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = [("순위", "rank", 6), ("종목명", "name", 24), ("코드", "code", 12),
               ("현재가", "price", 14), ("구분", "session", 8), ("등락", "direction", 8),
               ("변동", "change", 12), ("등락률", "rate", 10),
               ("거래대금", "trading_value", 16), ("링크", "url", 50)]
    wb = Workbook()
    wb.remove(wb.active)
    for key, items in data.items():
        if not items:
            continue
        ws = wb.create_sheet(titles[key][:31])  # 시트 이름은 31자 제한
        ws.append([h for h, _, _ in headers])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F5597")
            c.alignment = Alignment(horizontal="center")
        for it in items:
            ws.append([it[k] for _, k, _ in headers])
            link = ws.cell(row=ws.max_row, column=len(headers))
            link.hyperlink = link.value
            link.font = Font(color="0563C1", underline="single")
        for i, (_, _, w) in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="네이버 증권 랭킹 크롤러")
    parser.add_argument("-n", "--num", type=int, default=20, help="수집 개수 (API 방식, 기본 20)")
    parser.add_argument("-o", "--order", default="priceTop", choices=list(ORDER_TYPES),
                        help="정렬 기준 (기본 priceTop=거래대금 상위)")
    parser.add_argument("-m", "--market", default="ALL", choices=["ALL", "KOSPI", "KOSDAQ", "KONEX"],
                        help="시장 (기본 ALL)")
    parser.add_argument("--html", help="API 대신, 저장된 HTML 파일을 파싱")
    parser.add_argument("-g", "--group", choices=[v[0] for v in GROUPS.values()],
                        help="--html 사용 시 특정 그룹만 (kr/us/etf/crypto)")
    parser.add_argument("-j", "--json", help="JSON 저장 경로")
    parser.add_argument("-x", "--excel", help="엑셀(.xlsx) 저장 경로")
    args = parser.parse_args()

    if args.html:
        data = parse_rankings(read_html_file(args.html))
        titles = {k: f"{v[1]} 거래대금 상위" for k, v in GROUPS.items()}
        if args.group:
            data = {k: v for k, v in data.items() if GROUPS[k][0] == args.group}
        if not data:
            raise SystemExit(
                "HTML에서 랭킹 데이터(GlobalStockRankingV2)를 찾지 못했습니다.\n"
                "랭킹이 보이는 화면에서 개발자도구(F12) > Elements > <html> 우클릭 >\n"
                "Copy > Copy outerHTML 로 저장한 파일인지 확인하세요.")
    else:
        if args.group:
            raise SystemExit("-g 옵션은 --html과 함께만 사용할 수 있습니다. (API 방식은 국내주식만 지원)")
        titles = {"api": f"국내주식 {ORDER_TYPES[args.order]} ({args.market})"}
        try:
            print("수집 중...")
            items = fetch_domestic(
                args.order, args.market, limit=args.num,
                on_page=lambda p, got, total: print(f"  페이지 {p}: {got}건 (누적 {total}건)"))
        except requests.RequestException as e:
            raise SystemExit(f"수집 실패: {e}")
        except (ValueError, KeyError) as e:
            raise SystemExit(f"수집 실패: {e}")
        if not items:
            raise SystemExit("조회된 데이터가 없습니다.")
        data = {"api": items}

    print_rankings(data, titles)
    try:
        if args.json:
            with open(args.json, "w", encoding="utf-8") as f:
                json.dump({titles[k]: v for k, v in data.items()}, f, ensure_ascii=False, indent=2)
            print(f"\nJSON 저장: {args.json}")
        if args.excel:
            save_excel(data, titles, args.excel)
            print(f"엑셀 저장: {args.excel}")
    except PermissionError:
        raise SystemExit("저장 실패: 파일에 쓸 수 없습니다. 같은 이름의 파일이 열려 있다면 닫고 다시 시도하세요.")
    except OSError as e:
        raise SystemExit(f"저장 실패: {e}")


if __name__ == "__main__":
    main()
