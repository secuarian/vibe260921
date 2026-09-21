"""네이버 검색 결과(뉴스)에서 기사 목록을 파싱하고 본문까지 크롤링한다.

사용법:
    python news_crawler.py                 # 기본 URL, 기사 5건 (목록 + 본문)
    python news_crawler.py -n 10           # 기사 10건
    python news_crawler.py -q 인공지능     # 검색어 변경
    python news_crawler.py --list-only     # 검색 결과 목록만 (본문 요청 안 함)
    python news_crawler.py -o news.json    # 결과를 JSON 파일로 저장
    python news_crawler.py -x news.xlsx    # 결과를 엑셀 파일로 저장
필요 패키지: requests, beautifulsoup4, openpyxl(엑셀 저장 시)
"""
import argparse
import json
import time

import requests
from bs4 import BeautifulSoup

SEARCH_URL = (
    "https://search.naver.com/search.naver?where=nexearch&sm=top_hty"
    "&fbm=0&ie=utf8&query=%EB%B0%98%EB%8F%84%EC%B2%B4&ackey=ym8fy6yf"
)
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9",
}
DELAY = 1.0  # 서버 부담을 줄이기 위한 요청 간 대기(초)


def fetch(url, params=None):
    resp = requests.get(url, headers=HEADERS, params=params, timeout=10)
    resp.raise_for_status()
    resp.encoding = "utf-8"
    return BeautifulSoup(resp.text, "html.parser")


def clean_text(el):
    """요소의 텍스트를 반환. 숨김 안내문('새 창 열림')을 제거하고,
    <mark> 등 인라인 태그 앞뒤의 공백은 그대로 유지한다."""
    if el is None:
        return ""
    el = BeautifulSoup(str(el), "html.parser")
    for hidden in el.select(".fender-ui_0cb57fb2"):
        hidden.decompose()
    return " ".join(el.get_text().split())


def parse_search_results(soup, limit):
    """검색 결과 목록에서 기사 정보를 추출.

    각 기사 제목 링크는 data-heatmap-target=".tit" 속성을 가진다.
    제목 링크에서 위로 올라가며 언론사/시간/네이버뉴스 링크가 있는 블록을 찾는다.
    (메인 기사와 '관련 기사' 모두 같은 방식으로 처리)
    """
    items, seen = [], set()
    for title_a in soup.select('a[data-heatmap-target=".tit"]'):
        url = title_a.get("href", "")
        if not url or url in seen:
            continue
        seen.add(url)

        block = title_a.parent
        while block is not None and not block.select_one(
            ".sds-comps-profile-info-title-text"
        ):
            block = block.parent
        if block is None:
            continue

        press = block.select_one(".sds-comps-profile-info-title-text")
        subtext = block.select_one(".sds-comps-profile-info-subtext")
        nav = block.select_one('a[data-heatmap-target=".nav"]')
        body = block.select_one('a[data-heatmap-target=".body"]')

        items.append({
            "title": clean_text(title_a),
            "press": clean_text(press),
            "time": clean_text(subtext),
            "url": url,
            "naver_url": nav["href"] if nav else "",
            "snippet": clean_text(body),
        })
        if len(items) >= limit:
            break
    return items


def parse_article(url):
    """네이버 뉴스 기사 페이지에서 작성일과 본문을 추출."""
    soup = fetch(url)
    body = soup.select_one("#dic_area")
    date = soup.select_one(".media_end_head_info_datestamp_time")

    if body:
        # 본문 안의 사진 설명 등 불필요한 요소 제거
        for tag in body.select("script, style, .img_desc, .end_photo_org"):
            tag.decompose()

    return {
        "date": date.get("data-date-time", "") if date else "",
        "content": body.get_text("\n", strip=True) if body else "",
    }


EXCEL_COLUMNS = [
    # (헤더, 기사 dict 키, 열 너비)
    ("번호", None, 6),
    ("제목", "title", 50),
    ("언론사", "press", 14),
    ("시간", "time", 12),
    ("작성일", "date", 20),
    ("원문 URL", "url", 40),
    ("네이버 URL", "naver_url", 40),
    ("요약", "snippet", 50),
    ("본문", "content", 80),
]
EXCEL_CELL_LIMIT = 32767  # 엑셀 셀 하나에 넣을 수 있는 최대 글자 수


def save_excel(articles, path):
    """기사 목록을 엑셀(.xlsx) 파일로 저장."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"

    ws.append([name for name, _, _ in EXCEL_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, art in enumerate(articles, 1):
        row = [i]
        for _, key, _ in EXCEL_COLUMNS[1:]:
            value = ILLEGAL_CHARACTERS_RE.sub("", art.get(key, "") or "")
            row.append(value[:EXCEL_CELL_LIMIT])
        ws.append(row)
        for col, (_, key, _) in enumerate(EXCEL_COLUMNS, 1):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.alignment = Alignment(vertical="top", wrap_text=key in ("title", "snippet", "content"))
            if key in ("url", "naver_url") and cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")

    for col, (_, _, width) in enumerate(EXCEL_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="네이버 뉴스 크롤러")
    parser.add_argument("-q", "--query", help="검색어 (생략하면 기본 URL 사용)")
    parser.add_argument("-n", "--num", type=int, default=5, help="수집할 기사 수")
    parser.add_argument("--list-only", action="store_true",
                        help="검색 결과 목록만 수집 (기사 본문은 요청하지 않음)")
    parser.add_argument("-o", "--output", help="JSON 저장 경로")
    parser.add_argument("-x", "--excel", help="엑셀(.xlsx) 저장 경로")
    args = parser.parse_args()

    try:
        if args.query:
            soup = fetch("https://search.naver.com/search.naver",
                         {"where": "news", "query": args.query})
        else:
            soup = fetch(SEARCH_URL)
    except requests.RequestException as e:
        raise SystemExit(f"검색 페이지 요청 실패: {e}")

    articles = parse_search_results(soup, args.num)
    if not articles:
        raise SystemExit("기사를 찾지 못했습니다. 검색어를 바꾸거나 페이지 구조가 바뀌었는지 확인하세요.")
    print(f"기사 {len(articles)}건 발견\n")

    for i, art in enumerate(articles, 1):
        if not args.list_only and art["naver_url"]:
            try:
                art.update(parse_article(art["naver_url"]))
            except requests.RequestException as e:
                print(f"    (본문 수집 실패: {e})")
            time.sleep(DELAY)

        print(f"[{i}] {art['title']}")
        print(f"    {art['press']} | {art['time']}")
        print(f"    원문: {art['url']}")
        if art["naver_url"]:
            print(f"    네이버: {art['naver_url']}")
        text = art.get("content") or art["snippet"]
        if text:
            print(f"    {text[:200].replace(chr(10), ' ')}...")
        print()

    try:
        if args.output:
            with open(args.output, "w", encoding="utf-8") as f:
                json.dump(articles, f, ensure_ascii=False, indent=2)
            print(f"{len(articles)}건을 {args.output}에 저장했습니다.")

        if args.excel:
            save_excel(articles, args.excel)
            print(f"{len(articles)}건을 {args.excel}에 저장했습니다.")
    except PermissionError:
        raise SystemExit("저장 실패: 파일에 쓸 수 없습니다. 같은 이름의 파일이 열려 있다면 닫고 다시 시도하세요.")
    except OSError as e:
        raise SystemExit(f"저장 실패: {e}")


if __name__ == "__main__":
    main()
