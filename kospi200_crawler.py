"""네이버 증권에서 코스피200 지수 데이터를 수집한다.

https://stock.naver.com/market/stock/kr 은 자바스크립트로 화면을 그리는 페이지라
HTML을 BeautifulSoup으로 파싱하면 데이터가 비어 있다(표 헤더만 존재).
화면이 내부적으로 호출하는 JSON 주소(m.stock.naver.com/api)에서 같은 데이터를 가져온다.

사용법:
    python kospi200_crawler.py                # 현재 지수 + 구성종목 전체(페이징) + 최근 10일 시세
    python kospi200_crawler.py -d 30          # 최근 30일 시세
    python kospi200_crawler.py -x kospi200.xlsx   # 엑셀 저장
    python kospi200_crawler.py -o kospi200.json   # JSON 저장
필요 패키지: requests (엑셀 저장 시 openpyxl)
"""
import argparse
import json
import time

import requests

BASE = "https://m.stock.naver.com/api/index/KPI200"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
}


def get_json(path, params=None):
    resp = requests.get(BASE + path, headers=HEADERS, params=params, timeout=10)
    resp.raise_for_status()
    return resp.json()


PAGE_SIZE = 50   # 서버 허용 최대값 (100 이상은 400 오류)
MAX_PAGES = 20   # 응답이 이상할 때 무한 반복 방지
PAGE_DELAY = 0.3  # 페이지 사이 대기(초)


def fetch_all_stocks(on_page=None):
    """구성종목을 page=1,2,3... 으로 넘기며 빈 페이지가 나올 때까지 모두 수집."""
    stocks, seen = [], set()
    for page in range(1, MAX_PAGES + 1):
        rows = get_json("/enrollStocks", {"page": page, "pageSize": PAGE_SIZE})
        if not rows:
            break
        new = [r for r in rows if r["itemCode"] not in seen]
        seen.update(r["itemCode"] for r in new)
        stocks.extend(new)
        if on_page:
            on_page(page, len(rows), len(stocks))
        if len(rows) < PAGE_SIZE or not new:  # 마지막 페이지
            break
        time.sleep(PAGE_DELAY)
    return stocks


def collect(days, on_page=None):
    basic = get_json("/basic")
    info = get_json("/integration")
    prices = get_json("/price", {"page": 1, "pageSize": days})

    return {
        "summary": {
            "name": basic["stockName"],
            "close": basic["closePrice"],
            "change": basic["compareToPreviousClosePrice"],
            "direction": basic["compareToPreviousPrice"]["text"],
            "rate": basic["fluctuationsRatio"],
            "market_status": basic["marketStatus"],
            "traded_at": basic["localTradedAt"],
            "details": {i["key"]: i["value"] for i in info["totalInfos"]},
            "investors": info["dealTrendInfo"],       # 개인/외국인/기관 순매수(억)
            "updown": info["upDownStockInfo"],        # 상승/하락/보합 종목 수
        },
        "stocks": [
            {
                "rank": i,
                "code": s["itemCode"],
                "name": s["stockName"],
                "close": s["closePrice"],
                "change": s["compareToPreviousClosePrice"],
                "direction": s["compareToPreviousPrice"]["text"],
                "rate": s["fluctuationsRatio"],
                "volume": s["accumulatedTradingVolume"],
                "trading_value": s["accumulatedTradingValue"],  # 백만원
                "market_value": s["marketValue"],               # 억원
                "url": s["endUrl"],
            }
            for i, s in enumerate(fetch_all_stocks(on_page), 1)
        ],
        "prices": [
            {
                "date": p["localTradedAt"],
                "close": p["closePrice"],
                "change": p["compareToPreviousClosePrice"],
                "rate": p["fluctuationsRatio"],
                "open": p["openPrice"],
                "high": p["highPrice"],
                "low": p["lowPrice"],
            }
            for p in prices
        ],
    }


def print_result(data):
    s = data["summary"]
    print(f"[{s['name']}] {s['close']}  {s['direction']} {s['change']} ({s['rate']}%)")
    print(f"  기준: {s['traded_at']}  ({s['market_status']})")
    print("  " + " | ".join(f"{k} {v}" for k, v in s["details"].items()))
    inv = s["investors"]
    print(f"  순매수(억) 개인 {inv['personalValue']} / 외국인 {inv['foreignValue']}"
          f" / 기관 {inv['institutionalValue']}")
    u = s["updown"]
    print(f"  상승 {u['riseCount']} / 보합 {u['steadyCount']} / 하락 {u['fallCount']}\n")

    print(f"[구성종목 {len(data['stocks'])}개]  (시가총액순)")
    for t in data["stocks"]:
        print(f"  {t['rank']:>3}. {t['name']:<14} {t['close']:>10} {t['direction']} {t['change']:>8}"
              f" ({t['rate']}%)  시총 {t['market_value']}억")

    print("\n[일별 시세]")
    print(f"  {'날짜':<11}{'종가':>10}{'전일비':>9}{'등락률':>8}{'시가':>10}{'고가':>10}{'저가':>10}")
    for p in data["prices"]:
        print(f"  {p['date']:<11}{p['close']:>10}{p['change']:>9}{p['rate']:>7}%"
              f"{p['open']:>10}{p['high']:>10}{p['low']:>10}")


def save_excel(data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def sheet(ws, headers, rows):
        ws.append([h for h, _ in headers])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F5597")
            c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([r.get(k, "") for _, k in headers])
        for i, (h, _) in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) * 2 + 4)
        ws.freeze_panes = "A2"

    wb = Workbook()
    ws = wb.active
    ws.title = "일별시세"
    sheet(ws, [("날짜", "date"), ("종가", "close"), ("전일비", "change"),
               ("등락률(%)", "rate"), ("시가", "open"), ("고가", "high"), ("저가", "low")],
          data["prices"])
    sheet(wb.create_sheet("구성종목"),
          [("순위", "rank"), ("종목코드", "code"), ("종목명", "name"), ("현재가", "close"),
           ("전일비", "change"), ("등락", "direction"), ("등락률(%)", "rate"),
           ("거래량", "volume"), ("거래대금(백만)", "trading_value"),
           ("시가총액(억)", "market_value")],
          data["stocks"])
    s = data["summary"]
    rows = [{"k": "지수", "v": s["name"]}, {"k": "현재가", "v": s["close"]},
            {"k": "전일비", "v": f"{s['direction']} {s['change']}"},
            {"k": "등락률(%)", "v": s["rate"]}, {"k": "기준시각", "v": s["traded_at"]}]
    rows += [{"k": k, "v": v} for k, v in s["details"].items()]
    sheet(wb.create_sheet("요약"), [("항목", "k"), ("값", "v")], rows)
    wb.move_sheet("요약", offset=-2)
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="코스피200 크롤러")
    parser.add_argument("-d", "--days", type=int, default=10, help="일별 시세 일수")
    parser.add_argument("-o", "--output", help="JSON 저장 경로")
    parser.add_argument("-x", "--excel", help="엑셀(.xlsx) 저장 경로")
    args = parser.parse_args()

    def progress(page, got, total):
        print(f"  페이지 {page}: {got}건 수집 (누적 {total}건)")

    try:
        print("구성종목 수집 중...")
        data = collect(args.days, progress)
    except (requests.RequestException, KeyError, ValueError) as e:
        raise SystemExit(f"수집 실패: {e}")

    print_result(data)
    try:
        if args.output:
            with open(args.output, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            print(f"\nJSON 저장: {args.output}")
        if args.excel:
            save_excel(data, args.excel)
            print(f"엑셀 저장: {args.excel}")
    except PermissionError:
        raise SystemExit("저장 실패: 파일에 쓸 수 없습니다. 같은 이름의 파일이 열려 있다면 닫고 다시 시도하세요.")
    except OSError as e:
        raise SystemExit(f"저장 실패: {e}")


if __name__ == "__main__":
    main()
